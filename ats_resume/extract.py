"""Extract raw text from PDF, DOCX, or plain-text resumes."""

from pathlib import Path


def extract_text(path):
    """Return the plain text of a resume file (.pdf, .docx, .txt, .md)."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _from_pdf(path)
    if ext == ".docx":
        return _from_docx(path)
    if ext in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="ignore")
    raise ValueError(
        f"Unsupported file type '{ext}'. Please provide a PDF, DOCX, or TXT file."
    )


def _from_pdf(path):
    text = ""
    try:
        import pdfplumber

        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    except ImportError:
        text = ""
    if text.strip():
        return text
    # Fallback extractor
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ImportError as exc:  # pragma: no cover
        raise ImportError(
            "Reading PDFs requires 'pdfplumber' (recommended) or 'pypdf'.\n"
            "Install with:  pip install pdfplumber"
        ) from exc


def _from_docx(path):
    from docx import Document

    doc = Document(str(path))
    lines = []
    for p in doc.paragraphs:
        text = p.text
        if text.strip() and _is_list_paragraph(p):
            text = "\u2022 " + text.lstrip()
        lines.append(text)
    # Pull any text trapped inside tables so nothing is lost.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append("  ".join(cells))
    return "\n".join(lines)


def _is_list_paragraph(paragraph):
    """True if a DOCX paragraph is a bullet/numbered list item (so we can
    re-add a bullet marker that plain-text extraction would otherwise drop)."""
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "list" in style or "bullet" in style:
        return True
    pPr = paragraph._p.pPr
    return pPr is not None and pPr.numPr is not None


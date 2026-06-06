"""Render a structured resume dict into a clean, single-column, ATS-friendly DOCX.

Design choices that keep the output ATS-parseable:
  * single column, no tables / text boxes / images / headers-footers
  * standard, recognisable section headings
  * real bullet lists and plain text (selectable, machine-readable)
  * a common font (Calibri) and modest sizing so it fits on one page
"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Calibri"
# Headings whose contents are best shown as a single comma-joined line.
_INLINE_HEADINGS = {"Languages", "Interests"}
# Headings rendered as one flowing paragraph rather than bullets.
_PARAGRAPH_HEADINGS = {"Professional Summary"}


def _set_margins(doc, inches=0.5):
    for section in doc.sections:
        section.top_margin = Inches(inches)
        section.bottom_margin = Inches(inches)
        section.left_margin = Inches(inches)
        section.right_margin = Inches(inches)


def _bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _style_run(run, size=10, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    _style_run(p.add_run(text.upper()), size=11, bold=True)
    _bottom_border(p)
    return p


def _line(doc, text, size=10, bold=False, align=None, space_after=1, justify=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align:
        p.alignment = align
    _style_run(p.add_run(text), size=size, bold=bold)
    return p


def _bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.2)
    if bold_prefix:
        _style_run(p.add_run(bold_prefix), size=10, bold=True)
    _style_run(p.add_run(text), size=10)
    return p


def _two_col(doc, left, right, left_bold=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.5), WD_ALIGN_PARAGRAPH.RIGHT)
    _style_run(p.add_run(left), size=10, bold=left_bold)
    _style_run(p.add_run("\t" + right), size=10, bold=True)
    return p


def _contact_line(contact):
    order = ["location", "phone", "email", "linkedin", "github", "website"]
    parts = [contact[k] for k in order if contact.get(k)]
    return "  |  ".join(parts)


def _render_skills(doc, entries):
    for entry in entries:
        items = ([entry["title"]] if entry["title"] else []) + entry["bullets"]
        for item in items:
            if ":" in item:
                label, _, value = item.partition(":")
                _bullet(doc, value.strip(), bold_prefix=label.strip() + ": ")
            else:
                _bullet(doc, item)


def _render_inline(doc, entries):
    items = []
    for entry in entries:
        if entry["title"]:
            items.append(entry["title"])
        items.extend(entry["bullets"])
    if items:
        _line(doc, ", ".join(items))


def _render_paragraph(doc, entries):
    chunks = []
    for entry in entries:
        if entry["title"]:
            chunks.append(entry["title"])
        chunks.extend(entry["bullets"])
    text = " ".join(chunks).strip()
    if text:
        _line(doc, text, justify=True, space_after=2)


def _render_default(doc, entries):
    for entry in entries:
        title, date, bullets = entry["title"], entry["date"], entry["bullets"]
        if title and date:
            _two_col(doc, title, date, left_bold=True)
        elif title:
            _line(doc, title, bold=True, space_after=1)
        for b in bullets:
            _bullet(doc, b)


def build_docx(structured, output_path):
    """Build the resume DOCX from a structured dict and save it to ``output_path``."""
    doc = Document()
    _set_margins(doc, 0.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)

    # Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    _style_run(p.add_run(structured.get("name", "Your Name").upper()), size=18, bold=True)

    # Contact
    contact_line = _contact_line(structured.get("contact", {}))
    if contact_line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        _style_run(p.add_run(contact_line), size=10)

    for section in structured.get("sections", []):
        heading = section["heading"]
        entries = section.get("entries", [])
        if not entries:
            continue
        _heading(doc, heading)
        if heading == "Technical Skills":
            _render_skills(doc, entries)
        elif heading in _INLINE_HEADINGS:
            _render_inline(doc, entries)
        elif heading in _PARAGRAPH_HEADINGS:
            _render_paragraph(doc, entries)
        else:
            _render_default(doc, entries)

    doc.save(str(output_path))
    return output_path
